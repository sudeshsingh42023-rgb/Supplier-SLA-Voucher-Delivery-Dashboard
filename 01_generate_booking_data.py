import random
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

random.seed(42)

suppliers = ["SkyWings Consolidator", "GlobalStay Hotels", "HolidayCraft DMC",
             "AeroLink Fares", "UrbanNest Hospitality", "TerraTrip Packages"]
products = ["Flight", "Hotel", "Holiday Package"]
agents = [f"AGT-{1000+i}" for i in range(40)]
statuses_pool = ["Confirmed", "Confirmed", "Confirmed", "Delayed", "Cancelled", "Pending"]

start_date = datetime(2026, 6, 1)
rows = []
booking_id = 5001
for i in range(220):
    b_date = start_date + timedelta(days=random.randint(0, 89))
    supplier = random.choice(suppliers)
    product = random.choice(products)
    agent = random.choice(agents)
    gmv = round(random.uniform(8000, 185000), 0) if product != "Flight" else round(random.uniform(4000, 45000), 0)

    # Supplier-specific behavior baked in (so the dashboard finds real patterns)
    base_tat = {"SkyWings Consolidator": 4, "GlobalStay Hotels": 9, "HolidayCraft DMC": 14,
                "AeroLink Fares": 3, "UrbanNest Hospitality": 11, "TerraTrip Packages": 16}[supplier]
    tat_hours = max(1, round(random.gauss(base_tat, base_tat * 0.4), 1))
    sla_hours = {"Flight": 6, "Hotel": 24, "Holiday Package": 48}[product]
    voucher_status = "On Time" if tat_hours <= sla_hours else "Breach"

    payment_delay_days = max(0, round(random.gauss(2 if supplier != "TerraTrip Packages" and supplier != "HolidayCraft DMC" else 6, 2)))
    booking_status = random.choices(statuses_pool, weights=[55, 55, 55, 12, 8, 6])[0]

    kyc_status = random.choices(["Verified", "Pending", "Expired"], weights=[70, 22, 8])[0]

    rows.append({
        "Booking ID": f"TC{booking_id}",
        "Booking Date": b_date.strftime("%Y-%m-%d"),
        "Agent ID": agent,
        "Agent KYC Status": kyc_status,
        "Product": product,
        "Supplier": supplier,
        "GMV (INR)": gmv,
        "Booking Status": booking_status,
        "Voucher TAT (hrs)": tat_hours,
        "SLA Threshold (hrs)": sla_hours,
        "Voucher SLA Status": voucher_status,
        "Supplier Payment Delay (days)": payment_delay_days,
    })
    booking_id += 1

wb = openpyxl.Workbook()

# ---------- Sheet 1: Raw Data ----------
ws = wb.active
ws.title = "Raw Bookings"
headers = list(rows[0].keys())
ws.append(headers)
for r in rows:
    ws.append(list(r.values()))

header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(bold=True, color="FFFFFF", name="Arial")
for c in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=c)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center")

for c in range(1, len(headers) + 1):
    col_letter = get_column_letter(c)
    max_len = max(len(str(headers[c-1])), *(len(str(r[headers[c-1]])) for r in rows))
    ws.column_dimensions[col_letter].width = min(max_len + 3, 26)

for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for cell in row:
        cell.font = Font(name="Arial", size=10)

n = len(rows)
last_row = n + 1

wb.save("/home/claude/travclan/TravClan_Booking_Ops_SLA_Tracker.xlsx")
print("raw data written:", n, "rows")
